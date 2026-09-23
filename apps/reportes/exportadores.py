"""Exportacion de reportes a CSV y Excel.

La arquitectura esta preparada para PDF: basta agregar un exportador que
reciba (reporte, filas) y devuelva un HttpResponse.
"""
import csv
import datetime as dt
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse


def _valor_plano(valor):
    if isinstance(valor, (dt.date, dt.datetime)):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


def nombre_archivo(reporte, extension: str) -> str:
    marca = dt.datetime.now().strftime("%Y%m%d_%H%M")
    return f"{reporte.clave}_{marca}.{extension}"


def exportar_csv(reporte, filas) -> HttpResponse:
    respuesta = HttpResponse(content_type="text/csv; charset=utf-8")
    respuesta["Content-Disposition"] = (
        f'attachment; filename="{nombre_archivo(reporte, "csv")}"')
    respuesta.write("﻿")  # BOM para que Excel reconozca los acentos
    escritor = csv.writer(respuesta)
    escritor.writerow([etiqueta for _, etiqueta in reporte.columnas])
    for fila in filas:
        escritor.writerow([_valor_plano(fila.get(clave, "")) for clave, _ in reporte.columnas])
    return respuesta


def exportar_excel(reporte, filas) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte.clave[:31]

    hoja["A1"] = reporte.nombre
    hoja["A1"].font = Font(size=14, bold=True)
    hoja["A2"] = reporte.descripcion
    hoja["A2"].font = Font(size=9, italic=True)
    hoja["A3"] = f"Generado el {dt.datetime.now():%d/%m/%Y %H:%M} - Moneda: MXN"
    hoja["A3"].font = Font(size=9)

    fila_encabezado = 5
    relleno = PatternFill("solid", fgColor="1F3A5F")
    for indice, (_, etiqueta) in enumerate(reporte.columnas, start=1):
        celda = hoja.cell(row=fila_encabezado, column=indice, value=etiqueta)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")

    numero_fila = fila_encabezado + 1
    for fila in filas:
        for indice, (clave, _) in enumerate(reporte.columnas, start=1):
            valor = fila.get(clave, "")
            celda = hoja.cell(row=numero_fila, column=indice, value=_valor_plano(valor))
            if isinstance(valor, Decimal):
                celda.number_format = '"$"#,##0.00'
        numero_fila += 1

    for indice, (_clave, etiqueta) in enumerate(reporte.columnas, start=1):
        ancho = max(len(etiqueta) + 4, 14)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    respuesta = HttpResponse(
        buffer.read(),
        content_type=("application/vnd.openxmlformats-officedocument."
                      "spreadsheetml.sheet"))
    respuesta["Content-Disposition"] = (
        f'attachment; filename="{nombre_archivo(reporte, "xlsx")}"')
    return respuesta


def exportar_pdf(reporte, filas) -> HttpResponse:
    """Punto de extension para PDF (pendiente de implementar)."""
    raise NotImplementedError(
        "La exportacion a PDF aun no esta implementada. La arquitectura ya "
        "permite agregarla aqui sin tocar las vistas ni los generadores."
    )


EXPORTADORES = {"csv": exportar_csv, "excel": exportar_excel, "pdf": exportar_pdf}
