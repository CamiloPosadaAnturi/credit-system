"""Report export to CSV and Excel.

The architecture is ready for PDF: add an exporter that takes
(report, rows) and returns an HttpResponse.
"""
import csv
import datetime as dt
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils.translation import gettext as _


def _flat_value(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, Decimal):
        return float(value)
    return value


def file_name(report, extension: str) -> str:
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    return f"{report.key}_{stamp}.{extension}"


def export_csv(report, rows) -> HttpResponse:
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        f'attachment; filename="{file_name(report, "csv")}"')
    response.write("﻿")  # BOM so Excel detects the encoding
    writer = csv.writer(response)
    writer.writerow([str(label) for _key, label in report.columns])
    for row in rows:
        writer.writerow([_flat_value(row.get(key, "")) for key, _label in report.columns])
    return response


def export_excel(report, rows) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.key[:31]

    sheet["A1"] = str(report.name)
    sheet["A1"].font = Font(size=14, bold=True)
    sheet["A2"] = str(report.description)
    sheet["A2"].font = Font(size=9, italic=True)
    sheet["A3"] = _("Generated on %(stamp)s - Currency: MXN") % {
        "stamp": dt.datetime.now().strftime("%d/%m/%Y %H:%M")}
    sheet["A3"].font = Font(size=9)

    header_row = 5
    fill = PatternFill("solid", fgColor="1F3A5F")
    for index, (_key, label) in enumerate(report.columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=str(label))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    row_number = header_row + 1
    for row in rows:
        for index, (key, _label) in enumerate(report.columns, start=1):
            value = row.get(key, "")
            cell = sheet.cell(row=row_number, column=index, value=_flat_value(value))
            if isinstance(value, Decimal):
                cell.number_format = '"$"#,##0.00'
        row_number += 1

    for index, (_key, label) in enumerate(report.columns, start=1):
        width = max(len(str(label)) + 4, 14)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type=("application/vnd.openxmlformats-officedocument."
                      "spreadsheetml.sheet"))
    response["Content-Disposition"] = (
        f'attachment; filename="{file_name(report, "xlsx")}"')
    return response


def export_pdf(report, rows) -> HttpResponse:
    """Extension point for PDF (not implemented yet)."""
    raise NotImplementedError(
        _("PDF export is not implemented yet. The architecture already allows "
          "adding it here without touching the views or the generators.")
    )


EXPORTERS = {"csv": export_csv, "excel": export_excel, "pdf": export_pdf}
